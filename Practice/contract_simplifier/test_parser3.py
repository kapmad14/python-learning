# utils/parser.py
"""
Parser utilities: PDF/DOCX/text extraction and OCR via OCR.Space cloud.
Safe, non-raising OCR helper and clean PDF->text flow.
Public functions used by the app:
- extract_text_from_pdf(file)
- extract_text_from_docx(file)
- extract_text_from_image(file)
- extract_text_from_scanned_pdf(file)
"""

import io
import logging
from typing import Optional
import os

import pdfplumber
import requests
from docx import Document
import streamlit as st

logger = logging.getLogger(__name__)

# ---------- OCR.Space request helper (safe) ----------
def ocr_space_request(file_bytes: bytes, filename: str, language: str = "eng") -> str:
    """
    Send bytes to OCR.Space and return recognized text.
    This function is conservative: on any failure it logs and returns an empty string
    rather than raising an exception, so the app can show a friendly message.
    """
    # Lazy read the secret to avoid import-time errors
    OCR_KEY = st.secrets.get("OCR_SPACE_API_KEY", "") if hasattr(st, "secrets") else os.environ.get("OCR_SPACE_API_KEY", "")

    if not OCR_KEY:
        logger.warning("OCR_SPACE_API_KEY not set; OCR will be skipped and return empty text.")
        return ""

    url = "https://api.ocr.space/parse/image"
    payload = {
        "apikey": OCR_KEY,
        "language": language,
        "isOverlayRequired": False,
    }
    files = {"file": (filename, file_bytes)}

    try:
        resp = requests.post(url, data=payload, files=files, timeout=120)
        resp.raise_for_status()
    except requests.RequestException as e:
        logger.exception("OCR.Space request failed: %s", e)
        return ""

    try:
        result = resp.json()
    except Exception as e:
        logger.exception("OCR.Space returned non-JSON response: %s", e)
        return ""

    if result.get("IsErroredOnProcessing", False):
        # Log the error message(s) and return empty string
        err = result.get("ErrorMessage")
        logger.error("OCR.Space processing error: %s", err)
        return ""

    parsed_texts = []
    for pr in result.get("ParsedResults", []):
        parsed_texts.append(pr.get("ParsedText", ""))

    return "\n".join(parsed_texts).strip()

# ---------- pdfplumber-only extraction helper (internal) ----------
def _extract_text_from_pdf_plumber(file) -> str:
    """
    Use pdfplumber to extract selectable text.
    Accepts file-like (UploadedFile/BytesIO) or path.
    Returns extracted text or empty string on failure/no text.
    """
    try:
        # Prepare a stream for pdfplumber
        if hasattr(file, "getvalue"):
            b = file.getvalue()
            stream = io.BytesIO(b)
        elif isinstance(file, (bytes, bytearray)):
            stream = io.BytesIO(file)
        else:
            # assume file is a path-like object
            stream = file

        with pdfplumber.open(stream) as pdf:
            pages = []
            for page in pdf.pages:
                try:
                    t = page.extract_text() or ""
                    pages.append(t)
                except Exception as e:
                    logger.exception("pdfplumber page.extract_text failed on a page: %s", e)
                    pages.append("")
        return "\n".join(pages).strip()
    except Exception as e:
        logger.exception("pdfplumber extraction failed: %s", e)
        return ""

# ---------- Public PDF extraction (text-first, then OCR fallback) ----------
def extract_text_from_pdf(file) -> str:
    """
    Public helper used by the app.
    Tries pdfplumber first (fast, no external calls). If no selectable text found,
    falls back to sending the PDF bytes to OCR.Space (cloud OCR).
    Returns extracted text (possibly empty string on failure).
    """
    try:
        # Try pdfplumber selectable-text extraction first
        text = _extract_text_from_pdf_plumber(file)
        if text and text.strip():
            return text

        # No selectable text found -> fallback to OCR on PDF bytes
        # Prepare bytes and filename
        if hasattr(file, "getvalue"):
            b = file.getvalue()
            filename = getattr(file, "name", "file.pdf")
        elif isinstance(file, (bytes, bytearray)):
            b = file
            filename = "file.pdf"
        else:
            # file is path-like
            try:
                with open(file, "rb") as fh:
                    b = fh.read()
            except Exception as e:
                logger.exception("Failed to read file path for OCR fallback: %s", e)
                return ""
            filename = os.path.basename(file)

        if not filename.lower().endswith(".pdf"):
            filename = filename + ".pdf"

        text = ocr_space_request(b, filename)
        return text
    except Exception as e:
        logger.exception("extract_text_from_pdf encountered an unexpected error: %s", e)
        return ""

# ---------- DOCX extraction ----------
def extract_text_from_docx(file) -> str:
    """
    Extract text from a DOCX file. Accepts file-like or path.
    """
    try:
        if hasattr(file, "getvalue"):
            b = file.getvalue()
            stream = io.BytesIO(b)
            doc = Document(stream)
        else:
            doc = Document(file)
        paragraphs = [p.text for p in doc.paragraphs]
        return "\n".join(paragraphs).strip()
    except Exception as e:
        logger.exception("extract_text_from_docx failed: %s", e)
        return ""

# ---------- Image extraction (PNG/JPEG) ----------
def extract_text_from_image(file) -> str:
    """
    Send an image (uploaded file or bytes) to OCR.Space and return text.
    Safe: returns empty string on failures.
    """
    try:
        if hasattr(file, "getvalue"):
            b = file.getvalue()
            filename = getattr(file, "name", "image.png")
        elif isinstance(file, (bytes, bytearray)):
            b = file
            filename = "image.png"
        else:
            logger.error("extract_text_from_image: unsupported input type")
            return ""
        text = ocr_space_request(b, filename)
        return text
    except Exception as e:
        logger.exception("extract_text_from_image failed: %s", e)
        return ""

# ---------- Scanned PDF extraction (explicit) ----------
def extract_text_from_scanned_pdf(file) -> str:
    """
    Explicitly send PDF bytes to OCR.Space. Accepts UploadedFile, bytes, or path.
    """
    try:
        if hasattr(file, "getvalue"):
            b = file.getvalue()
            filename = getattr(file, "name", "file.pdf")
        elif isinstance(file, (bytes, bytearray)):
            b = file
            filename = "file.pdf"
        else:
            try:
                with open(file, "rb") as fh:
                    b = fh.read()
            except Exception as e:
                logger.exception("extract_text_from_scanned_pdf failed to read file path: %s", e)
                return ""
            filename = os.path.basename(file)

        if not filename.lower().endswith(".pdf"):
            filename = filename + ".pdf"

        text = ocr_space_request(b, filename)
        return text
    except Exception as e:
        logger.exception("extract_text_from_scanned_pdf failed: %s", e)
        return ""
